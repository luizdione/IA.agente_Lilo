"""DOCX: leitura em texto e escrita simples a partir de Markdown (grupo analise).

Escrita cobre o que um relatorio de analise precisa: titulos (#, ##, ###),
paragrafos, listas (-, *, 1.) e tabelas em pipe (| a | b |). Nada de estilos
finos: para manuscrito de verdade, o caminho e a skill build_paper.
"""
from __future__ import annotations

import re
from pathlib import Path

from .comum import erro, resolver


def ler(caminho: str) -> str:
    """Texto de um .docx: paragrafos (com marca de titulo) e tabelas em pipe."""
    p = resolver(caminho)
    if not p.is_file():
        return erro("arquivo nao existe: %s" % p)
    if p.suffix.lower() != ".docx":
        return erro("%s nao e .docx" % p.name)
    import docx  # python-docx
    try:
        d = docx.Document(str(p))
    except Exception as e:  # noqa: BLE001
        return erro("DOCX ilegivel: %s" % e)
    partes: list[str] = []
    for par in d.paragraphs:
        t = par.text.strip()
        if not t:
            continue
        estilo = (par.style.name or "").lower() if par.style is not None else ""
        m = re.match(r"heading (\d)", estilo)
        if m:
            partes.append("#" * int(m.group(1)) + " " + t)
        elif estilo.startswith("title"):
            partes.append("# " + t)
        elif "list" in estilo:
            partes.append("- " + t)
        else:
            partes.append(t)
    for i, tab in enumerate(d.tables, 1):
        linhas = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in tab.rows]
        if not linhas:
            continue
        partes.append("\n[tabela %d]" % i)
        partes.append("| " + " | ".join(linhas[0]) + " |")
        partes.append("|" + "---|" * len(linhas[0]))
        for row in linhas[1:]:
            partes.append("| " + " | ".join(row) + " |")
    return "\n".join(partes) if partes else "(documento vazio)"


def _tabela(doc, linhas: list[list[str]]) -> None:
    ncol = max(len(r) for r in linhas)
    t = doc.add_table(rows=len(linhas), cols=ncol)
    t.style = "Table Grid"
    for i, row in enumerate(linhas):
        for j, cel in enumerate(row):
            t.cell(i, j).text = cel
            if i == 0:
                for run in t.cell(i, j).paragraphs[0].runs:
                    run.bold = True


def escrever(caminho: str, markdown: str, titulo: str = "") -> str:
    """Grava um .docx a partir de texto Markdown simples. Sobrescreve se existir."""
    p = resolver(caminho)
    if p.suffix.lower() != ".docx":
        return erro("o caminho precisa terminar em .docx")
    import docx
    doc = docx.Document()
    if titulo:
        doc.add_heading(titulo, level=0)
    tabela: list[list[str]] = []
    paragrafo: list[str] = []

    def fecha_paragrafo():
        if paragrafo:
            doc.add_paragraph(" ".join(paragrafo))
            paragrafo.clear()

    def fecha_tabela():
        if tabela:
            _tabela(doc, [r for r in tabela if not all(re.fullmatch(r":?-+:?", c) for c in r)])
            tabela.clear()

    for linha in markdown.splitlines():
        s = linha.rstrip()
        if s.startswith("|") and s.endswith("|"):
            fecha_paragrafo()
            tabela.append([c.strip() for c in s.strip("|").split("|")])
            continue
        fecha_tabela()
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            fecha_paragrafo()
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", s):
            fecha_paragrafo()
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", s), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", s):
            fecha_paragrafo()
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", s), style="List Number")
        elif not s.strip():
            fecha_paragrafo()
        else:
            paragrafo.append(s.strip())
    fecha_paragrafo()
    fecha_tabela()
    p.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(p))
    except OSError as e:
        return erro("gravacao: %s" % e)
    return "gravado: %s (%d bytes)" % (p, p.stat().st_size)
