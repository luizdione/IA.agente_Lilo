"""
agents/docgen/agent.py — Subagente de GERAÇÃO DE DOCUMENTOS do LILO (V13.2).
=========================================================================
Gera arquivos de escritório a partir de conteúdo/dados:
  - DOCX: Markdown -> Word (via pandoc; robusto, preserva tabelas/headings).
  - XLSX: CSV/TSV/dados -> planilha formatada (openpyxl: header em negrito,
    autofilter, largura de coluna).
  - PPTX: outline -> slides (python-pptx).
  - report: o LLM (cloud, LILO_DOCGEN_MODEL) redige um relatório em Markdown a
    partir do tema e entrega um .docx pronto.

O LLM gera o CONTEÚDO; as bibliotecas montam o arquivo (determinístico). Não
substitui o subagente `paper` (manuscritos cientificos) — docgen e para
relatorios/planilhas/apresentacoes do dia a dia.
"""

import csv
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from shared.base_agent import BaseAgent
from shared.llm_client import get_client
from shared.utils import load_env, sanitize_query, truncate
from shared.logger import get_logger

load_env()
log = get_logger("DocGenAgent")

DOCGEN_MODEL = os.getenv("LILO_DOCGEN_MODEL", "gpt-oss:120b-cloud")
_OUT_DIR = Path(os.getenv("LILO_DOCGEN_OUT", str(_ROOT / "data" / "docgen")))


def _out_path(name: str, ext: str) -> Path:
    _OUT_DIR.mkdir(parents=True, exist_ok=True)
    safe = "".join(c for c in name if c.isalnum() or c in " _-").strip()[:60]
    return _OUT_DIR / ("%s.%s" % (safe or "documento", ext))


class DocGenAgent(BaseAgent):
    """Gera DOCX/XLSX/PPTX a partir de conteudo, dados ou tema."""

    name = "docgen"
    description = (
        "Geracao de documentos de escritorio: relatorios Word (.docx), planilhas "
        "Excel (.xlsx) e apresentacoes (.pptx) a partir de conteudo, dados ou "
        "tema. LLM gera o conteudo; bibliotecas montam o arquivo."
    )
    KEYWORDS = [
        "gerar documento", "criar documento", "relatorio word", "docx",
        "planilha", "xlsx", "excel", "apresentacao", "slides", "pptx",
        "powerpoint", "gerar relatorio", "exportar para word",
    ]
    cost_tier = "ollama"
    local_only = False
    latency_class = "fast"

    def __init__(self):
        super().__init__()
        self._llm = None

    @property
    def llm(self):
        if self._llm is None:
            self._llm = get_client()
        return self._llm

    # ── DOCX (markdown -> Word via pandoc) ───────────────────────────────────
    def _md_to_docx(self, md: str, out: Path) -> str:
        with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False,
                                         encoding="utf-8") as fh:
            fh.write(md)
            src = fh.name
        try:
            subprocess.run(["pandoc", src, "-o", str(out)],
                           capture_output=True, text=True, timeout=120, check=True)
            return "DOCX gerado: %s" % out
        except Exception as e:
            return "[docgen] falha no pandoc: %s" % e
        finally:
            try:
                os.unlink(src)
            except OSError:
                pass

    # ── XLSX (CSV/TSV/linhas -> planilha formatada) ──────────────────────────
    def _to_xlsx(self, rows: list[list], out: Path, title: str = "Dados") -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            wb = Workbook()
            ws = wb.active
            ws.title = title[:31] or "Dados"
            for r in rows:
                ws.append(list(r))
            if rows:
                for c in ws[1]:
                    c.font = Font(bold=True)
                ws.auto_filter.ref = ws.dimensions
                for col in ws.columns:
                    w = max((len(str(c.value)) for c in col if c.value), default=10)
                    ws.column_dimensions[col[0].column_letter].width = min(60, w + 2)
            wb.save(out)
            return "XLSX gerado: %s (%d linhas)" % (out, len(rows))
        except Exception as e:
            return "[docgen] falha no xlsx: %s" % e

    @staticmethod
    def _parse_rows(ctx: dict) -> list[list]:
        if ctx.get("csv_path") and Path(ctx["csv_path"]).exists():
            with open(ctx["csv_path"], newline="", encoding="utf-8") as fh:
                return [row for row in csv.reader(fh)]
        raw = ctx.get("tsv") or ctx.get("csv") or ""
        if raw:
            delim = "\t" if "\t" in raw else ","
            return [line.split(delim) for line in raw.strip().splitlines()]
        return []

    # ── PPTX (outline -> slides) ─────────────────────────────────────────────
    def _to_pptx(self, outline: str, out: Path) -> str:
        try:
            from pptx import Presentation
            from pptx.util import Pt
            prs = Presentation()
            bullet = prs.slide_layouts[1]   # layout "Title and Content"
            cur_bullets, cur_title = None, None

            def _flush(title, bullets):
                if title is None:
                    return
                s = prs.slides.add_slide(bullet)
                s.shapes.title.text = title
                body = s.placeholders[1].text_frame
                body.clear()
                for i, b in enumerate(bullets):
                    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                    p.text = b
                    p.font.size = Pt(18)

            for line in outline.splitlines():
                ln = line.rstrip()
                if not ln:
                    continue
                if ln.startswith("#"):
                    _flush(cur_title, cur_bullets or [])
                    cur_title = ln.lstrip("# ").strip()
                    cur_bullets = []
                elif ln.lstrip().startswith(("-", "*", "•")):
                    (cur_bullets if cur_bullets is not None else []).append(
                        ln.lstrip(" -*•"))
                else:
                    if cur_bullets is not None:
                        cur_bullets.append(ln.strip())
            _flush(cur_title, cur_bullets or [])
            prs.save(out)
            return "PPTX gerado: %s (%d slides)" % (out, len(prs.slides._sldIdLst))
        except Exception as e:
            return "[docgen] falha no pptx: %s" % e

    # ── BaseAgent ─────────────────────────────────────────────────────────────
    def run(self, query: str, context: dict[str, Any] = {}) -> str:
        ctx = context or {}
        mode = ctx.get("mode", "report")
        name = ctx.get("name") or sanitize_query(query)[:50] or "documento"
        log.info("DocGen: mode=%s | %s", mode, truncate(query, 60))

        if mode == "xlsx":
            rows = self._parse_rows(ctx)
            if not rows:
                return "[docgen] sem dados (use csv_path/csv/tsv no contexto)"
            return self._to_xlsx(rows, _out_path(name, "xlsx"), title=name)

        if mode == "pptx":
            outline = ctx.get("outline") or query
            return self._to_pptx(outline, _out_path(name, "pptx"))

        if mode == "docx":
            md = ctx.get("content") or query
            return self._md_to_docx(md, _out_path(name, "docx"))

        # report: LLM (cloud) redige o relatorio em Markdown -> DOCX.
        sys_p = ("Voce e um analista que redige relatorios tecnicos claros e "
                 "objetivos em Markdown (titulos #, listas, tabelas pipe). "
                 "PT-BR. Nao invente dados; se faltarem, sinalize.")
        try:
            md = self.llm.chat(
                "Redija um relatorio em Markdown sobre:\n%s\n\n%s" % (
                    query, ("DADOS/CONTEXTO:\n" + ctx["data"]
                            if ctx.get("data") else "")),
                # V18 QW2: loga divergencia policy->agente (modelo inalterado).
                system=sys_p, model=self.resolve_model(ctx, DOCGEN_MODEL),
                temperature=0.3, max_tokens=4096)
        except Exception as e:
            return "[docgen] LLM (%s) falhou: %s" % (DOCGEN_MODEL, e)
        docx_msg = self._md_to_docx(md, _out_path(name, "docx"))
        return "## Relatorio gerado\n\n%s\n\n---\n%s" % (docx_msg, md[:1500])

    def get_tools(self) -> list[dict]:
        return [
            {"name": "make_report_docx",
             "description": "LLM redige um relatorio sobre o tema e entrega um .docx.",
             "inputSchema": {"type": "object", "properties": {
                 "query": {"type": "string"}, "name": {"type": "string"},
                 "data": {"type": "string"}}, "required": ["query"]}},
            {"name": "make_docx",
             "description": "Converte Markdown/conteudo em .docx (pandoc).",
             "inputSchema": {"type": "object", "properties": {
                 "content": {"type": "string"}, "name": {"type": "string"}},
                 "required": ["content"]}},
            {"name": "make_xlsx",
             "description": "Gera planilha .xlsx de CSV/TSV (csv_path, csv ou tsv).",
             "inputSchema": {"type": "object", "properties": {
                 "csv_path": {"type": "string"}, "tsv": {"type": "string"},
                 "name": {"type": "string"}}}},
            {"name": "make_pptx",
             "description": "Gera apresentacao .pptx de um outline (# titulo / - bullets).",
             "inputSchema": {"type": "object", "properties": {
                 "outline": {"type": "string"}, "name": {"type": "string"}},
                 "required": ["outline"]}},
        ]

    def capabilities(self) -> dict:
        return {"name": self.name, "keywords": self.KEYWORDS,
                "model": DOCGEN_MODEL, "out_dir": str(_OUT_DIR),
                "formats": ["docx", "xlsx", "pptx"]}

    def health(self) -> dict:
        def _has(mod):
            try:
                __import__(mod)
                return True
            except Exception:
                return False
        import shutil
        return {"status": "ok", "model": DOCGEN_MODEL,
                "pandoc": bool(shutil.which("pandoc")),
                "python_docx": _has("docx"), "openpyxl": _has("openpyxl"),
                "python_pptx": _has("pptx")}
